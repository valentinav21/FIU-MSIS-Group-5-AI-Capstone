"""Excel workbook with every result, formatted for buyers and sales managers."""
from __future__ import annotations

import io as _io

import numpy as np
import pandas as pd

MONEY_HINTS = ("value", "revenue", "spend", "cost", "price", "total", "gap", "leakage", "opportunity", "margin_12m", "expected_value")
PCT_HINTS = ("rate", "share", "%", "probability", "growth", "service", "target", "risk", "pct", "W1", "W2", "W3", "W4", "W5", "W6", "W7", "W8")


def _fmt_sheet(ws, df):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    head = PatternFill("solid", fgColor="1F3A5F")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.freeze_panes = "B2" if df.shape[1] > 3 else "A2"
    ws.row_dimensions[1].height = 32
    for j, col in enumerate(df.columns, start=1):
        name = str(col)
        low = name.lower()
        width = min(max(len(name), 10) + 2, 60)
        if df.shape[0]:
            sample = df[col].astype(str).head(200).str.len()
            width = int(min(max(width, sample.quantile(0.9) + 2 if len(sample) else width), 70))
        ws.column_dimensions[get_column_letter(j)].width = width
        fmt = None
        if pd.api.types.is_numeric_dtype(df[col]):
            if any(h in low for h in ("pct", "rate", "share", "probability", "growth", "service", "target")) or name in [f"W{k}" for k in range(1, 9)] or "%" in name:
                fmt = "0.0%"
            elif any(h in low for h in MONEY_HINTS):
                fmt = '"$"#,##0'
            else:
                fmt = "#,##0.##"
        if fmt:
            for cell in ws.iter_cols(min_col=j, max_col=j, min_row=2, max_row=ws.max_row):
                for c in cell:
                    c.number_format = fmt
        if low in ("why", "root_causes", "detail", "suggested_action", "alternatives", "what_it_means", "reading"):
            for cell in ws.iter_cols(min_col=j, max_col=j, min_row=2, max_row=min(ws.max_row, 3000)):
                for c in cell:
                    c.alignment = Alignment(wrap_text=False)
    ws.auto_filter.ref = ws.dimensions


def _clean(df):
    df = df.copy()
    for c in df.columns:
        if pd.api.types.is_datetime64_any_dtype(df[c]):
            df[c] = df[c].dt.date
        elif df[c].dtype == object:
            df[c] = df[c].map(lambda v: ", ".join(map(str, v)) if isinstance(v, (list, tuple, set)) else v)
    return df.replace([np.inf, -np.inf], np.nan)


def workbook(res) -> bytes:
    sheets = []
    summ = [("Created", res.created), ("Data folder", res.data_folder), ("Data as of", str(res.as_of.date()) if res.as_of is not None else ""),
            ("Planning bucket", "Weekly" if res.bucket == "W" else "Monthly")]
    ps = res.policy_summary or {}
    if ps:
        summ += [("Items x warehouses planned", ps.get("planned_series")),
                 ("Target fill rate (revenue weighted)", ps.get("fill_rate")),
                 ("Today's SAP settings fill rate (simulated)", ps.get("current_fill_rate")),
                 ("Safety stock value, AI plan", ps.get("safety_stock_value")),
                 ("Safety stock value, today's SAP minimums", ps.get("current_safety_stock_value"))]
    if res.budget.get("same_as_today"):
        summ.append(("Fill rate with today's safety stock budget, AI allocation", res.budget["same_as_today"]["fill_rate"]))
    if len(res.po_recs):
        summ += [("PO lines recommended", len(res.po_recs)), ("PO value recommended", float(res.po_recs["line_total"].sum()))]
    if len(res.alerts):
        for k, v in res.alerts["risk"].value_counts().items():
            summ.append((f"Alerts: {k}", int(v)))
    for n in res.notes:
        summ.append(("Note", n))
    for e in res.errors:
        summ.append(("Error", e))
    sheets.append(("Summary", pd.DataFrame(summ, columns=["Item", "Value"])))
    if len(res.po_recs):
        sheets.append(("PO recommendations", res.po_recs))
    if len(res.transfers):
        sheets.append(("Transfers", res.transfers))
    if len(res.alerts):
        sheets.append(("Risk alerts", res.alerts))
    if len(res.policy):
        cols = [c for c in res.policy.columns]
        sheets.append(("Inventory policy", res.policy[cols]))
    if len(res.forecast_table):
        sheets.append(("Forecast", res.forecast_table))
        sheets.append(("Forecast accuracy", res.forecast_accuracy))
    if len(res.calibration):
        sheets.append(("Forecast calibration", res.calibration))
    if len(res.supplier_scorecard):
        sheets.append(("Supplier scorecard", res.supplier_scorecard))
    if res.replay.get("table") is not None:
        sheets.append(("Policy replay", res.replay["table"].reset_index().rename(columns={"index": "Measure"})))
        sheets.append(("Model check", res.replay["model_check"]))
    S = res.sales
    if S is not None:
        for name, df in (("Sales plays", S.plays), ("Customers", S.customers), ("Segments", S.segments), ("Personas", S.personas),
                         ("Cross-sell", S.recommendations), ("Bundles", S.bundles), ("Upsell gaps", S.upsell),
                         ("Geography", S.geo), ("Price leakage", S.pricing_leakage), ("Price sensitivity", S.elasticity),
                         ("Discounts by size", S.discount_check), ("Products", S.products), ("Categories", S.categories),
                         ("Unmet web demand", S.unmet_demand), ("Searched but out of stock", S.search_vs_stock)):
            if df is not None and len(df):
                sheets.append((name, df))
    sheets.append(("Data health", res.health))
    sheets.append(("Column mapping", res.mapping))
    buf = _io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        for name, df in sheets:
            d = _clean(df)
            if len(d) > 200000:
                d = d.head(200000)
            d.to_excel(xw, sheet_name=name[:31], index=False)
            _fmt_sheet(xw.sheets[name[:31]], d)
    return buf.getvalue()
